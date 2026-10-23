"""Exports the sheet "Monster" of the user's workbook docs/RPG_0.1.xlsx: every number and ability of the 91 monsters.

The sheet replaces what src/data/raw/MONSTER_DATA.json (from the old g.py) says about a monster; src/data/tables.ts lays the two together.
Column "entity 9: outside_spawner_hut_xp": a number = the monster has a Monster Hut on the overworld (the number is the XP of burning it
down), "x" = the monster lives on the overworld without a hut, empty = it does not appear outside.

Run:  python -I tools/export/export_monster_sheet.py   (or: npm run data:monsters)
Out:  src/data/monster-sheet.json
"""

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SHEET = ROOT / "docs" / "RPG_0.1.xlsx"
OUT = ROOT / "src" / "data" / "monster-sheet.json"

# header text (part of it) -> key; the sheet names both petrified and death "entity 8i", the status word tells them apart
ABILITY_COLUMNS = [
    ("ranged attack", "ranged"),
    ("hits_all", "hitsAll"),
    ("causes_status_poisoned", "poisoned"),
    ("causes_status_cursed", "cursed"),
    ("causes_status_sleep", "sleeping"),
    ("causes_status_insanity", "crazy"),
    ("causes_status_disease", "diseased"),
    ("causes_status_weakness", "weak"),
    ("causes_status_petrified", "petrified"),
    ("causes_status_death", "dead"),
    ("causes_status_paralysis", "paralyzed"),
    ("causes_status_eradication", "eradication"),
]


def find(head, text):
    for i, h in enumerate(head):
        if h and text in str(h).lower():
            return i
    raise KeyError(text)


def number(value):
    if value in (None, ""):
        return None
    return float(value) if isinstance(value, float) and not float(value).is_integer() else int(value)


def main() -> int:
    ws = openpyxl.load_workbook(SHEET, data_only=True)["Monster"]
    head = [c.value for c in ws[1]]
    col = {
        "name": find(head, "entity 0"),
        "health": find(head, "entity 1:"),
        "armor": find(head, "entity 2:"),
        "speed": find(head, "entity 3:"),
        "attacks": find(head, "entity 4:"),
        "damageMin": find(head, "damage_min"),
        "damageMax": find(head, "damage_max"),
        "attackType": find(head, "attack_type"),
        "regen": find(head, "regeneration"),
        "stealMana": find(head, "steal_mana"),
        "aging": find(head, "increases_magical_age"),
        "hut": find(head, "entity 9"),
        "resMagic": find(head, "res_magic"),
        "resFire": find(head, "res_fire"),
        "resPhysical": find(head, "res_physical"),
        "resCold": find(head, "res_cold"),
        "resPoison": find(head, "res_poison"),
        "gold": find(head, "entity 16"),
        "iconHeight": find(head, "icon_height"),
        "terrain": find(head, "entity 19"),
        "undead": find(head, "entity 20"),
        "boss": find(head, "final_boss"),
        "group": find(head, "entity 22"),
        "xp": find(head, "entity 23:"),
        "level": find(head, "entity 24:"),
    }
    abilities = {key: find(head, text) for text, key in ABILITY_COLUMNS}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        m = {}
        for key, i in col.items():
            v = row[i]
            m[key] = v.strip() if isinstance(v, str) else number(v)
        m["undead"] = str(m["undead"]).lower() == "true"
        hut = m["hut"]
        m["hut"] = "x" if isinstance(hut, str) and hut.lower() == "x" else hut
        for key, i in abilities.items():
            m[key] = isinstance(row[i], str) and row[i].strip().lower() == "x"
        age = row[col["aging"]]
        m["aging"] = 1 if isinstance(age, str) and age.strip().lower() == "x" else number(age)
        out[str(int(row[0]))] = m
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"monsters: {len(out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
