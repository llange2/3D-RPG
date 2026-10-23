"""Exports what the user defined in the documentation workbooks that the old g.py does not carry.

- old_code/documentation/MM3_Master_documentation.xlsx, sheet "Spells": name, description and explanation of all 78 spells.
  The numbers (mana, level, damage, school) of the sheet equal src/data/raw/SPELLS.json, only the texts differ (the new spells
  have their own names), so only the texts are exported.
- docs/RPG_0.1.xlsx (the user's current workbook), sheet "Monster": the column "entity 22: group" says who is undead (the flag of entity 20 is
  not used for that).

Run:  python -I tools/export/export_spell_sheet.py
Out:  src/data/spell-sheet.json, src/data/monster-groups.json
"""

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT.parent / "old_code" / "documentation"
OUT = ROOT / "src" / "data"


def clean(value):
    return value.strip() if isinstance(value, str) else value


def export_spells() -> int:
    ws = openpyxl.load_workbook(DOCS / "MM3_Master_documentation.xlsx", data_only=True)["Spells"]
    head = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(head, row))
        if d["ID"] is None:
            continue
        out[str(int(d["ID"]))] = {
            "name": clean(d["Name"]),
            "description": clean(d["Description"]),
            "explanation": clean(d.get("Explanation")) or "",
            "status": clean(d["Comment"]) or "",
        }
    (OUT / "spell-sheet.json").write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return len(out)


def export_groups() -> int:
    ws = openpyxl.load_workbook(ROOT / "docs" / "RPG_0.1.xlsx", data_only=True)["Monster"]
    head = [c.value for c in ws[1]]
    group_col = next(i for i, h in enumerate(head) if h and "group" in str(h))
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[group_col] is None:
            continue
        out[str(int(row[0]))] = clean(row[group_col])
    (OUT / "monster-groups.json").write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return len(out)


if __name__ == "__main__":
    print(f"spells: {export_spells()}, monster groups: {export_groups()}")
    sys.exit(0)
