"""
Reads the user's list of map objects (.ignore/items.xlsx) and writes src/data/item-rules.ts: which old items must always do something
(never pure decoration) and in which kinds of place each may stand. The generators and a test read the result.

Usage: python -I tools/export/export_item_rules.py
"""
import json
import os
import re

import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
SOURCE = os.path.join(ROOT, ".ignore", "items.xlsx")
OUT = os.path.join(ROOT, "src", "data", "item-rules.ts")

# column header -> environment id used by the game
ENVIRONMENTS = {"Castle": "castle", "Cave": "cave", "Dungeon": "dungeon", "Town": "town", "Spaceship": "spaceship", "Outside": "outside"}

# Corrections the user made to the list after the sheet was written (the sheet itself stays as it is): item -> environments to add.
# 118 Gong: "you can put gongs in caves - my mistake" (2026-10-20).
# 3 Portal: "every town must have a portal" (2026-10-20, nineteenth playtest, point 8).
# 3 Portal also in dungeons (quest document 1.2, answer to open point 1: "kann auch in Dungeons eingesetzt werden", it links two cells A>B>A and is
# separate from the town portals). 17..25 the keys also in caves: the end rooms of level 1 and 2 caves hold keys (document 1.2, "End-Raeume").
ADD_PLACES = {118: ["cave"], 3: ["town", "dungeon"], **{k: ["cave"] for k in range(17, 26)}}
# 73 Magic runes: a quest object (the way into the end room of quest_3, the barriers of dungeon_3), never plain decoration (open point 4).
ADD_FUNCTIONAL = {73}

DOC = os.path.join(ROOT, "docs", "OBJECTS.md")

wb = openpyxl.load_workbook(SOURCE)
ws = wb.active
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    label = row[0]
    if not label:
        continue
    m = re.match(r"ID(\d+):\s*(.*)", str(label))
    if not m:
        continue
    item = int(m.group(1))
    name = m.group(2).strip()
    functional = bool(row[1]) or item in ADD_FUNCTIONAL
    envs = [ENVIRONMENTS[header[i]] for i in range(2, 8) if row[i]]
    for extra in ADD_PLACES.get(item, []):
        if extra not in envs:
            envs.append(extra)
    comment = (row[8] or "").strip() if len(row) > 8 else ""
    rows.append((item, name, functional, envs, comment))

lines = [
    "/**",
    " * Which old map objects may stand where, and which must always do something. Written by tools/export/export_item_rules.py from the",
    " * user's list (.ignore/items.xlsx): do not edit by hand.",
    " *",
    " * `functional`: the object needs a use (a lever pulls something, a barrel can be searched, a chest holds loot), even if it is a dummy; it",
    " * is never placed as pure decoration. `places`: the kinds of place it may stand in.",
    " */",
    'export type Place = "castle" | "cave" | "dungeon" | "town" | "spaceship" | "outside";',
    "",
    "export interface ItemRule {",
    "  name: string;",
    "  functional: boolean;",
    "  places: readonly Place[];",
    "}",
    "",
    "export const ITEM_RULES: Readonly<Record<number, ItemRule>> = {",
]
for item, name, functional, envs, _comment in rows:
    places = ", ".join(json.dumps(e) for e in envs)
    lines.append("  %d: { name: %s, functional: %s, places: [%s] }," % (item, json.dumps(name), "true" if functional else "false", places))
lines += [
    "};",
    "",
    "/** May this object stand in this kind of place at all? Objects the list does not know are not restricted. */",
    "export const allowedIn = (item: number, place: Place): boolean => ITEM_RULES[item]?.places.includes(place) ?? true;",
    "",
    "/** Must this object always do something (no pure decoration)? */",
    "export const isFunctional = (item: number): boolean => ITEM_RULES[item]?.functional ?? false;",
    "",
]
with open(OUT, "w", encoding="utf-8", newline="\n") as f:
    f.write("\n".join(lines))
print("wrote %s: %d items" % (OUT, len(rows)))

# the same list as a table for people and agents: docs/OBJECTS.md
order = ["castle", "cave", "dungeon", "town", "spaceship", "outside"]
doc = [
    "# Objekte des Spiels: wo sie stehen dürfen und was sie tun müssen",
    "",
    "**Diese Liste gilt für alles, was ein Generator oder eine Regeldatei stellt** (Städte, Burgen, Höhlen, Dungeons, Pyramide/Raumschiff, Überwelt, Quests, Verkäufer).",
    "Vor jeder Änderung an der Generierung (Regeln `src/data/rules/*.json`, Themes `src/data/dungeon-themes.ts`, Quests `src/actions/questGen.ts`, Städte `src/interiors/town.ts`,",
    "Überwelt `src/world/*`, `src/data/talent-sellers.ts`) hier nachsehen, welches Objekt wo stehen darf.",
    "",
    "- **Funktional = ja:** das Objekt muss eine Verwendung haben (ein Hebel zieht etwas, ein Fass ist ein Behälter, eine Truhe gibt etwas, ein Tor öffnet sich), auch wenn es nur eine Attrappe ist. Es ist **nie reine Deko**.",
    "- **Orte:** ein Objekt steht nur in den Spalten, die für es gesetzt sind (Burg, Höhle, Dungeon, Stadt, Raumschiff = Innenraum der Pyramide, Überwelt).",
    "- Quelle ist die Liste des Nutzers `.ignore/items.xlsx` plus die Korrekturen im Kopf von `tools/export/export_item_rules.py` (`ADD_PLACES`). Daraus entstehen **generiert** diese Datei und `src/data/item-rules.ts`",
    "  (`python -I tools/export/export_item_rules.py`): nicht von Hand ändern. Änderungswünsche des Nutzers an der Liste kommen in die Excel-Datei oder nach `ADD_PLACES`, dann neu erzeugen.",
    "- Im Code: `allowedIn(item, place)`, `isFunctional(item)` (`item-rules.ts`), `placeOf(map)` (`src/interiors/place.ts`), `itemIssues` in `validateContent` (ein Test über alle Orte schlägt bei einem Verstoß fehl), `plainProps(theme)` filtert die Deko.",
    "- Jedes Objekt, das der Spieler sieht, sagt beim Betreten oder Anstoßen einen Satz (`src/data/prop-texts.ts`, `propLine`); neue Objekte bekommen eine Zeile. Bäume und Berge sind die Ausnahme.",
    "",
    "| ID | Objekt | Funktional | Burg | Höhle | Dungeon | Stadt | Raumschiff | Überwelt | Anmerkung |",
    "|---|---|---|---|---|---|---|---|---|---|",
]
for item, name, functional, envs, comment in rows:
    cells = ["x" if e in envs else "" for e in order]
    doc.append("| %d | %s | %s | %s | %s |" % (item, name, "**ja**" if functional else "", " | ".join(cells), comment.replace("|", "/")))
with open(DOC, "w", encoding="utf-8", newline="\n") as f:
    f.write("\n".join(doc) + "\n")
print("wrote %s" % DOC)
